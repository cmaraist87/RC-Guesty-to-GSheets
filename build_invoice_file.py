"""
Write the invoice to an .xlsx file, formulas and all.

    python build_invoice_file.py            # -> Invoice RCI-2026-01.xlsx

No credentials, no sharing, no setup. Open it in Excel, or drag it into Google
Drive and open it as a Google Sheet -- the formulas survive either way.

Everything that should calculate, calculates:

    Amount      = Hours x Rate, blank while the line is blank
    Subtotal    = SUM over the whole item block
    Discount    = a flat figure or a percentage, chosen from a dropdown
    Amount due  = Subtotal less the discount, never below zero

Twenty-five item rows are laid out and the SUM covers all of them, so a line can
be added by typing into a blank one and removed by deleting the row. Neither
touches a formula.
"""
from __future__ import annotations

import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT = "Invoice RCI-2026-01.xlsx"

HEAD = 15
FIRST = 16
N_ITEMS = 25
LAST = FIRST + N_ITEMS - 1      # 40
R_HOURS, R_SUB, R_DISC, R_DUE = LAST + 2, LAST + 3, LAST + 4, LAST + 5

FILL = "<< fill in >>"

INK = "131820"
FAINT = "8A94A1"
ACCENT = "17594E"
ACCENT_BG = "E8F1EE"
FLAG_BG = "FBF1DF"
FLAG_INK = "8A5A12"
HAIR = "EEF1F4"

ITEMS = [
    ("Discovery, Guesty API integration and data pipeline",
     "Authentication, paginated reservation retrieval, field mapping, and the transformation into schedule rows.", 12, 125),
    ("Scheduling logic",
     "Same-day turnover detection, early and late arrival/departure codes, and multi-unit properties split into one clean per unit.", 6.5, 125),
    ("Google Sheets write layer",
     "Writing without destroying the team's tick boxes or formatting; monthly tab routing and automatic creation of new months.", 7, 125),
    ("Daily change tracking and visual marking",
     "Highlighting of new and amended bookings, permanent strikethrough of cancellations, and per-cell colour coding for turnovers and adjusted times.", 5, 125),
    ("Reliability engineering and automated scheduling",
     "Shared credential caching within the provider's daily limit, once-per-day execution that tolerates late delivery, and safeguards that stop rather than corrupt the sheet on bad data.", 6, 125),
    ("Cloud infrastructure and deployment",
     "Google Cloud project, storage and service-account configuration; scheduled deployment pipeline with manual override and safety toggle.", 5, 125),
    ("Data remediation of existing records",
     "Removal of ~1,600 duplicated rows and 454 out-of-market rows; correction of cancellation marks that had drifted onto live bookings.", 4.5, 125),
    ("Testing, deployment supervision and documentation",
     "142 automated checks; supervised live runs across several mornings; written handover covering operation and the manual controls.", 6.5, 125),
]

MONEY = '"$"#,##0.00'


# A percentage rather than a flat figure, deliberately. "15% family discount"
# reads as a decision; "$984.38 off" reads as an arbitrary number somebody picked.
# It also stays right if a line is added or removed later.
#
# STANDING, not per-phase. The client's choice: this is a family relationship and
# the discount is meant to hold across future work, so the label says nothing about
# a phase. The Rate column still shows $125, so the standard rate stays on the
# record -- which is what a referral reads, and what the next quote is anchored to.
DISCOUNT_LABEL = "Family discount"
DISCOUNT_MODE = "%"
DISCOUNT_VALUE = 15


def build(path: str = OUT, label: str = DISCOUNT_LABEL,
          mode: str = DISCOUNT_MODE, value: float = DISCOUNT_VALUE) -> str:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"

    label_font = Font(name="Calibri", size=8, bold=True, color=FAINT)
    body = Font(name="Calibri", size=11, color=INK)
    small = Font(name="Calibri", size=9, color=FAINT)
    top = Alignment(vertical="top", wrap_text=True)
    right = Alignment(horizontal="right")
    hair = Border(bottom=Side(style="thin", color=HAIR))

    for col, w in zip("ABCDE", (46, 52, 9, 11, 14)):
        ws.column_dimensions[col].width = w

    def put(cell, value, font=None, align=None, fmt=None, border=None):
        c = ws[cell]
        c.value = value
        c.font = font or body
        if align: c.alignment = align
        if fmt: c.number_format = fmt
        if border: c.border = border
        return c

    put("A1", "INVOICE", Font(name="Calibri", size=26, bold=True, color=INK))
    put("A2", "Software development & systems integration", small)
    for i, (k, v) in enumerate([("Invoice No", "RCI-2026-01"), ("Issued", "2026-09-06"),
                                ("Due", "2026-09-21"), ("Terms", "Net 15")], start=1):
        put(f"D{i}", k, label_font, right)
        put(f"E{i}", v)

    put("A5", "FROM", label_font)
    put("A6", "Christopher Maraist", Font(name="Calibri", size=12, bold=True, color=INK))
    put("A7", "chris.tektonpro@gmail.com")
    put("A8", f"{FILL} street, city, state, ZIP")
    put("A9", f"{FILL} phone")

    put("C5", "BILL TO", label_font)
    put("C6", "Ramos Cleaning, Inc.", Font(name="Calibri", size=12, bold=True, color=INK))
    put("C7", f"{FILL} street, city, state, ZIP")
    put("C8", f"{FILL} accounts payable contact")

    put("A11", "PROJECT", label_font)
    put("A12", "Guesty → Google Sheets automation", Font(name="Calibri", size=11, bold=True, color=INK))
    put("A13", "Nightly cleaning-schedule build across five markets · 2026-08-02 → 2026-09-04", small)

    head_border = Border(bottom=Side(style="medium", color=INK))
    for col, txt, al in (("A", "Services rendered", None), ("B", "Detail", None),
                         ("C", "Hours", right), ("D", "Rate", right), ("E", "Amount", right)):
        put(f"{col}{HEAD}", txt, Font(name="Calibri", size=9, bold=True, color=FAINT),
            al, border=head_border)

    for i in range(N_ITEMS):
        r = FIRST + i
        if i < len(ITEMS):
            desc, detail, hours, rate = ITEMS[i]
            put(f"A{r}", desc, body, top, border=hair)
            put(f"B{r}", detail, small, top, border=hair)
            put(f"C{r}", hours, body, right, "0.##", hair)
            put(f"D{r}", rate, body, right, MONEY, hair)
        else:
            for col in "ABCD":
                put(f"{col}{r}", None, small if col == "B" else body, top, border=hair)
            ws[f"C{r}"].number_format = "0.##"
            ws[f"D{r}"].number_format = MONEY
        # Blank until the line has both figures, so the spare rows stay invisible.
        put(f"E{r}", f'=IF(OR($C{r}="",$D{r}=""),"",$C{r}*$D{r})', body, right, MONEY, hair)
        ws.row_dimensions[r].height = 30

    bold_r = Font(name="Calibri", size=11, bold=True, color=INK)
    put(f"D{R_HOURS}", "Total hours", bold_r, right)
    put(f"E{R_HOURS}", f"=SUM(C{FIRST}:C{LAST})", bold_r, right, "0.##")
    put(f"D{R_SUB}", "Subtotal", bold_r, right)
    put(f"E{R_SUB}", f"=SUM(E{FIRST}:E{LAST})", bold_r, right, MONEY)

    put(f"B{R_DISC}", label, bold_r, right)
    put(f"C{R_DISC}", mode, Font(name="Calibri", size=11, bold=True, color=ACCENT),
        Alignment(horizontal="center"))
    # A percentage is a plain number in this cell; a flat discount is money.
    put(f"D{R_DISC}", value, body, right, "0.##\"%\"" if mode == "%" else MONEY)
    # Negative, so the total is a plain addition and the deduction reads as one.
    put(f"E{R_DISC}", f'=IF($C{R_DISC}="%",-$E{R_SUB}*$D{R_DISC}/100,-$D{R_DISC})',
        body, right, MONEY)

    due_font = Font(name="Calibri", size=14, bold=True, color=ACCENT)
    due_fill = PatternFill("solid", fgColor=ACCENT_BG)
    due_border = Border(top=Side(style="thin", color=ACCENT),
                        bottom=Side(style="thin", color=ACCENT))
    for col, val, fmt in (("D", "AMOUNT DUE", None),
                          ("E", f"=MAX(0,$E{R_SUB}+$E{R_DISC})", MONEY)):
        c = put(f"{col}{R_DUE}", val, due_font, right, fmt, due_border)
        c.fill = due_fill

    put(f"A{R_DUE + 2}", "PAYMENT", label_font)
    for i, (k, v) in enumerate([("Payable to", "Christopher Maraist"),
                                ("Method", f"{FILL} bank transfer / check / Zelle"),
                                ("Reference", "RCI-2026-01")]):
        put(f"A{R_DUE + 3 + i}", k, small)
        put(f"B{R_DUE + 3 + i}", v)

    put(f"A{R_DUE + 7}", "NOT INCLUDED", label_font)
    n = put(f"A{R_DUE + 8}",
            "Work on the Connecteam job-scheduling integration is under way and is "
            "deliberately excluded here. It will be billed separately once jobs are live "
            "in the Connecteam account. Ongoing monitoring and maintenance is also not "
            "included and can be arranged under separate terms.", small, top)
    ws.merge_cells(f"A{R_DUE + 8}:E{R_DUE + 8}")
    ws.row_dimensions[R_DUE + 8].height = 46
    put(f"A{R_DUE + 10}", "Please include the invoice reference with payment.", small)

    # $ or % only: the discount formula branches on this cell, and "USD" would
    # quietly produce the wrong total rather than an error.
    dv = DataValidation(type="list", formula1='"$,%"', allow_blank=False, showDropDown=False)
    dv.error = "Choose $ for a flat amount or % for a percentage."
    dv.errorTitle = "Discount type"
    ws.add_data_validation(dv)
    dv.add(ws[f"C{R_DISC}"])

    # Make the placeholders impossible to miss.
    ws.conditional_formatting.add(
        f"A1:E{R_DUE + 6}",
        FormulaRule(formula=[f'ISNUMBER(SEARCH("{FILL}",A1))'],
                    fill=PatternFill("solid", fgColor=FLAG_BG),
                    font=Font(color=FLAG_INK)))

    ws.freeze_panes = f"A{HEAD + 1}"
    ws.sheet_view.showGridLines = False
    ws.print_area = f"A1:E{R_DUE + 10}"
    ws.page_setup.orientation = "portrait"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    wb.save(path)
    return path


def main(argv=None) -> int:
    import argparse

    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("out", nargs="?", default=OUT)
    ap.add_argument("--discount", default=f"{DISCOUNT_VALUE}%",
                    help='e.g. "15%%" for a percentage, or "1562.50" for a flat amount.')
    ap.add_argument("--discount-label", default=DISCOUNT_LABEL)
    args = ap.parse_args(argv)

    raw = str(args.discount).strip()
    mode = "%" if raw.endswith("%") else "$"
    try:
        value = float(raw.rstrip("%$").replace(",", "").strip() or 0)
    except ValueError:
        print(f"ERROR: could not read --discount {args.discount!r}", file=sys.stderr)
        return 2

    try:
        out = build(args.out, label=args.discount_label, mode=mode, value=value)
    except PermissionError:
        # Excel holds an exclusive lock on an open workbook. Say so, rather than
        # showing a zipfile traceback for what is really "close the file".
        print(f"ERROR: '{args.out}' is open in Excel, so it cannot be rewritten.",
              file=sys.stderr)
        print("       Close it and run this again, or pass a different filename.",
              file=sys.stderr)
        return 3
    sub_total = sum(h * r for _, _, h, r in ITEMS)
    off = sub_total * value / 100 if mode == "%" else value
    due = max(0.0, sub_total - off)
    print(f"Wrote {out}")
    print(f"  Subtotal            ${sub_total:,.2f}")
    print(f"  {args.discount_label:<18}  -${off:,.2f}  ({raw})")
    print(f"  AMOUNT DUE          ${due:,.2f}")
    print()
    print(f"  line items  rows {FIRST}-{LAST}  (type into a blank row to add, delete a row to remove)")
    print(f"  discount    row {R_DISC}: $ or % in column C, the figure in column D")
    print(f"  amount due  row {R_DUE}, calculated")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
